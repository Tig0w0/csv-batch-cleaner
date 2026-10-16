"""Two real US-government datasets with labeled synthetic data-quality anomalies.

Run once to download frozen source CSVs, and again to reproduce the exact fixtures.
Only data/test_cases/ is written; the original USGS fixture is untouched.
"""
import csv
import hashlib
import json
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1] / "data" / "test_cases"
NOAA_URL = "https://www.ncei.noaa.gov/access/services/data/v1?" + urlencode({
    "dataset": "daily-summaries", "stations": "USW00094728",
    "startDate": "2024-01-01", "endDate": "2024-05-31",
    "dataTypes": "TMAX,TMIN,PRCP", "units": "metric", "format": "csv",
    "includeAttributes": "false",
})
CENSUS_URL = ("https://www2.census.gov/programs-surveys/popest/datasets/"
              "2020-2024/state/totals/NST-EST2024-ALLDATA.csv")


def source_csv(agency, url):
    path = ROOT / agency / "original.csv"
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        with urlopen(Request(url, headers={"User-Agent": "Data-Quality-Portfolio/1.0"}), timeout=30) as response:
            contents = response.read()
        if not contents or len(contents) > 2_000_000:
            raise ValueError(f"{agency}: unexpectedly sized download")
        path.write_bytes(contents)
    with path.open(encoding="utf-8-sig", newline="") as stream:
        records = list(csv.DictReader(stream))
    print(f"{agency} source: {len(records)} rows, SHA256 {hashlib.sha256(path.read_bytes()).hexdigest()[:16]}")
    return records


def write_csv(path, fields, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(fields)
        writer.writerows(rows)


def prepare_noaa():
    raw = source_csv("noaa", NOAA_URL)
    assert len(raw) == 152, f"Expected 152 frozen NYC daily observations, found {len(raw)}"
    columns = ("date", "station", "prcp", "tmax", "tmin", "notes", "test_tag")
    base = [[r["DATE"], r["STATION"], r["PRCP"], r["TMAX"], r["TMIN"], "", "source"]
            for r in raw]
    assert len({r[0] for r in base}) == len(base), "NOAA daily date key must be unique"
    files = [list(base[:76]), list(base[76:])]
    # Four ID-level duplicates (one with conflicting numeric content).
    for index in range(4):
        clone = base[index].copy()
        clone[3] = "999.0" if index == 3 else clone[3]
        clone[-1] = "synthetic_duplicate_modified" if index == 3 else "synthetic_duplicate"
        files[1].append(clone)
    # Two valid *synthetic* records: optional blank number and formula-like note text.
    files[1].append(["2024-06-01", "USW00094728", "0.0", "20.0", "",
                     "Synthetic optional TMIN is empty", "synthetic_valid_optional_blank"])
    files[1].append(["2024-06-02", "USW00094728", "0.0", "20.0", "15.0",
                     '=HYPERLINK("https://example.invalid","unsafe")', "synthetic_formula_like_text"])
    # Eight synthetic invalid records exercising different validation paths.
    files[1].extend([
        ["", "USW00094728", "0", "10", "5", "", "synthetic_missing_date_a"],
        ["", "USW00094728", "0", "11", "6", "", "synthetic_missing_date_b"],
        ["2099-01-03", "USW00094728", "0", "not-a-number", "5", "", "synthetic_invalid_tmax_a"],
        ["2099-01-04", "USW00094728", "0", "infinity", "5", "", "synthetic_invalid_tmax_b"],
        ["2099-01-05", "USW00094728", "NaN", "11", "5", "", "synthetic_nonfinite_prcp"],
        ["2099-01-06", "", "0", "11", "5", "", "synthetic_missing_station"],
        ["2099-01-07", "USW00094728", "0", "11", "5", "", "synthetic_overflow", "EXTRA-VALUE"],
        ["2099-01-08", "USW00094728", "0", "11", "5", ""],  # too few cells
    ])
    for i, records in enumerate(files, 1):
        write_csv(ROOT / "noaa" / "input" / f"weather_part_{i:02d}.csv", columns, records)
    return {"source": 152, "injected_valid": 2, "injected_duplicates": 4,
            "injected_rejected": 8, "expected": {"files": 2, "rows": 166, "clean": 154, "duplicates": 4, "rejected": 8}}


def prepare_census():
    raw = source_csv("census", CENSUS_URL)
    states = [record for record in raw if record["SUMLEV"] == "040"]
    assert len(states) == 52, f"Expected 52 state-level geographies, got {len(states)}"
    columns = ("record_id", "state_code", "state_name", "year", "population",
               "region", "division", "test_tag")
    rows = []
    for entry in states:
        for year in (2021, 2022, 2023, 2024):
            rows.append([f'{entry["STATE"]}-{year}', int(entry["STATE"]), entry["NAME"],
                         year, int(entry[f"POPESTIMATE{year}"]),
                         entry["REGION"], entry["DIVISION"], "source"])
    assert len(rows) == 208 and len({r[0] for r in rows}) == 208
    original = [row.copy() for row in rows[:4]]
    for idx, row in enumerate(original):
        row[-1] = "synthetic_duplicate"
        if idx == 3:
            row[4] += 12345
            row[-1] = "synthetic_duplicate_modified"
    rows.extend(original)
    # Nine rejected synthetic cases: blank required fields, malformed numbers,
    # Excel error cell, spillover column and invalid year.
    errors = [
        ["", 99, "TEST", 2024, 1000, "9", "9", "synthetic_missing_id"],
        ["SYN-02", "", "TEST", 2024, 1000, "9", "9", "synthetic_missing_state_code"],
        ["SYN-03", 99, "TEST", 2024, "", "9", "9", "synthetic_missing_population"],
        ["SYN-04", 99, "TEST", 2024, "not-a-number", "9", "9", "synthetic_bad_population"],
        ["SYN-05", 99, "TEST", 2024, "NaN", "9", "9", "synthetic_nonfinite_population"],
        ["SYN-06", 99, "TEST", 2024, "#DIV/0!", "9", "9", "synthetic_excel_error"],
        ["SYN-07", 99, "TEST", 2024, 1000, "9", "9", "synthetic_extra_column", "SPILLOVER"],
        ["SYN-08", 99, "TEST", "", 1000, "9", "9", "synthetic_missing_year"],
        ["SYN-09", 99, "TEST", "20XX", 1000, "9", "9", "synthetic_invalid_year"],
    ]
    rows.extend(errors)
    file = ROOT / "census" / "input" / "state_population_with_issues.xlsx"
    file.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "State population"
    sheet.append(columns)
    for values in rows:
        sheet.append(values)
        cell = sheet.cell(sheet.max_row, 2)
        if isinstance(cell.value, int):
            cell.number_format = "00"  # Verify leading-zero state IDs stay intact.
        if values[0] == "SYN-06":
            sheet.cell(sheet.max_row, 5).data_type = "e"
    workbook.save(file)
    return {"source": 208, "injected_valid": 0, "injected_duplicates": 4,
            "injected_rejected": 9, "expected": {"files": 1, "rows": 221, "clean": 208, "duplicates": 4, "rejected": 9}}


def main():
    noaa = prepare_noaa()
    census = prepare_census()
    (ROOT / "manifest.json").write_text(json.dumps({
        "attribution": {"noaa": NOAA_URL, "census": CENSUS_URL},
        "source_sha256": {agency: hashlib.sha256((ROOT / agency / "original.csv").read_bytes()).hexdigest()
                          for agency in ("noaa", "census")},
        "notice": "Government source records are authentic. All test_tag=synthetic_* entries are invented for validation.",
        "scenarios": {"noaa_weather": noaa, "census_population_xlsx": census},
    }, indent=2) + "\n", encoding="utf-8")
    print("Created 2 additional independently testable tasks. Inputs and source snapshots are separate.")


if __name__ == "__main__":
    main()
