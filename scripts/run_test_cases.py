"""Exercise three independent data-cleaning tasks; never changes source records."""
import csv
import json
import sys
import tempfile
from dataclasses import asdict
from pathlib import Path

from openpyxl import Workbook, load_workbook

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT))
from cleaner import process


CASES = {
    "usgs_earthquakes": {
        "input": PROJECT / "data/input",
        "output": PROJECT / "data/test_cases/usgs/output",
        "key": "id", "required": "id,time,latitude,longitude",
        "numeric": "latitude,longitude,depth,magnitude",
        "expected": (3, 195, 180, 6, 9),
    },
    "noaa_weather": {
        "input": PROJECT / "data/test_cases/noaa/input",
        "output": PROJECT / "data/test_cases/noaa/output",
        "key": "date", "required": "date,station",
        "numeric": "prcp,tmax,tmin",
        "expected": (2, 166, 154, 4, 8),
    },
    "census_population_xlsx": {
        "input": PROJECT / "data/test_cases/census/input",
        "output": PROJECT / "data/test_cases/census/output",
        "key": "record_id", "required": "record_id,state_code,year,population",
        "numeric": "year,population",
        "expected": (1, 221, 208, 4, 9),
    },
}


def csv_rows(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def inspect(name, result):
    folder = Path(result.output_directory)
    cleaned, duplicates, rejected = [
        csv_rows(folder / f"{part}.csv") for part in ("cleaned", "duplicates", "rejected")
    ]
    assert (len(cleaned), len(duplicates), len(rejected)) == (
        result.clean_rows, result.duplicate_rows, result.rejected_rows
    )
    assert len(cleaned) + len(duplicates) + len(rejected) == result.processed_rows
    assert all(row["__source_file"] and row["__source_row"] for row in cleaned + duplicates + rejected)
    for category in ("cleaned", "duplicates", "rejected"):
        assert (folder / f"{category}.xlsx").is_file()
    if name == "noaa_weather":
        injected_accepted = {x["test_tag"] for x in cleaned if x["test_tag"].startswith("synthetic")}
        assert injected_accepted == {"synthetic_valid_optional_blank", "synthetic_formula_like_text"}
        unsafe = next(x for x in cleaned if x["test_tag"] == "synthetic_formula_like_text")
        assert unsafe["notes"].startswith("'=HYPERLINK"), unsafe["notes"]
        assert any(x["__extra_values"] == '["EXTRA-VALUE"]' for x in rejected)
        assert result.issues["malformed_csv_row"] == 2
        assert result.issues["duplicate_date"] == 4
    if name == "census_population_xlsx":
        assert any(x["state_code"] == "01" and x["state_name"] == "Alabama" for x in cleaned)
        book = load_workbook(folder / "cleaned.xlsx", read_only=True)
        try:
            items = list(book.active.values)
            state_idx = items[0].index("state_code")
            assert any(row[state_idx] == "01" for row in items[1:])
        finally:
            book.close()
        assert any(x["test_tag"] == "synthetic_excel_error" for x in rejected)
        assert any(x["__extra_values"] == '["SPILLOVER"]' for x in rejected)
        assert result.issues["excel_error_cell"] == 1
        assert result.issues["duplicate_record_id"] == 4
    if name == "usgs_earthquakes":
        assert result.issues["duplicate_id"] == 6
    return str(folder.parent.relative_to(PROJECT))  # Stable path; each run gets a fresh subfolder.


def expect_abort(title, make_input, required, numeric, message):
    with tempfile.TemporaryDirectory() as directory:
        source, out = Path(directory) / "input", Path(directory) / "output"
        source.mkdir()
        make_input(source)
        try:
            process(source, out, required=required, numeric=numeric)
        except (ValueError, OSError) as exc:
            assert message in str(exc), f"{title}: unexpected cause: {exc}"
            assert not any(out.glob("run_*")), f"{title}: partial results were published"
            return {"scenario": title, "status": "correctly_aborted", "error": str(exc)}
        raise AssertionError(f"{title}: should have failed but succeeded")


def extra_negative_cases():
    def malformed(directory):
        (directory / "broken.csv").write_text('id,name\nA,"unterminated\nB,broken\n', encoding="utf-8")
    def missing_schema(directory):
        with (directory / "wrong.csv").open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerows([["id", "amount"], ["A", "1"]])
    def formula(directory):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["id", "amount"])
        sheet.append(["A", "=1+2"])
        workbook.save(directory / "formula.xlsx")
    return [
        expect_abort("unclosed_csv_quote", malformed, "id", "", "malformed CSV"),
        expect_abort("configured_column_absent", missing_schema, "id", "price", "configured columns missing"),
        expect_abort("formula_in_xlsx", formula, "id", "amount", "formulas are not supported"),
    ]


def main():
    results = {}
    for name, cfg in CASES.items():
        summary = process(cfg["input"], cfg["output"], cfg["key"], cfg["required"], cfg["numeric"])
        counts = (summary.input_files, summary.processed_rows, summary.clean_rows,
                  summary.duplicate_rows, summary.rejected_rows)
        assert counts == cfg["expected"], f"{name}: expected {cfg['expected']}, got {counts}"
        folder = inspect(name, summary)
        results[name] = {"counts": counts, "output": folder, "issues": summary.issues}
        print(f"{name}: files={counts[0]}, rows={counts[1]}, clean={counts[2]}, "
              f"duplicates={counts[3]}, rejected={counts[4]} PASS")
    negative = extra_negative_cases()
    for item in negative:
        print(f"{item['scenario']}: rejected malformed input, no partial output PASS")
    totals = {key: sum(v["counts"][idx] for v in results.values())
              for idx, key in enumerate(("files", "rows", "clean", "duplicates", "rejected"))}
    print("TOTAL:", totals)
    report = {"results": results, "negative_tests": negative, "totals": totals,
              "scope": "Three independent tasks. Synthetic errors labeled; real source snapshots preserved."}
    report_path = PROJECT / "data/test_cases/test_results.json"
    report_path.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print("Result summary:", report_path)


if __name__ == "__main__":
    main()
