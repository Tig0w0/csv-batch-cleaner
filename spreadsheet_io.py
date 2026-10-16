"""Minimal CSV/XLSX input and XLSX output adapters.

The cleaning and validation rules live in cleaner.py, not here.
Only the first worksheet of each XLSX workbook is supported.
"""
from __future__ import annotations

import csv
import math
from datetime import date, datetime
from pathlib import Path
from typing import Iterator
from zipfile import BadZipFile


def _openpyxl():
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError as exc:
        raise RuntimeError("XLSX support requires: python -m pip install -r requirements.txt") from exc
    return Workbook, load_workbook


def read_header(path: Path) -> list[str]:
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            try:
                header = next(csv.reader(handle, strict=True), None)
            except csv.Error as exc:
                raise ValueError(f"{path.name}: invalid CSV header: {exc}") from exc
        return list(header or [])
    _, load_workbook = _openpyxl()
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            values = next(workbook.worksheets[0].iter_rows(values_only=True), None)
            return [str(value).strip() if value is not None else "" for value in (values or [])]
        finally:
            workbook.close()
    except (OSError, ValueError, IndexError, BadZipFile) as exc:
        raise ValueError(f"{path.name}: cannot read XLSX header: {exc}") from exc


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def read_records(path: Path, header_length: int) -> Iterator[tuple[int, list[str], bool]]:
    """Yield (physical source row, values, malformed).

    CSV source rows refer to the starting physical line, including quoted multiline fields.
    """
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle, strict=True)
            try:
                next(reader)
                previous_end = reader.line_num
                while True:
                    first_line = previous_end + 1
                    try:
                        values = next(reader)
                    except StopIteration:
                        break
                    previous_end = reader.line_num
                    yield first_line, [value.strip() for value in values[:header_length]], (
                        len(values) != header_length
                    )
            except csv.Error as exc:
                raise ValueError(f"{path.name}: malformed CSV near line {reader.line_num}: {exc}") from exc
        return

    _, load_workbook = _openpyxl()
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            rows = workbook.worksheets[0].iter_rows(values_only=False)
            next(rows, None)  # header
            for row_number, row in enumerate(rows, 2):
                if all(cell.value is None for cell in row):
                    continue
                if any(cell.data_type == "f" for cell in row):
                    raise ValueError(
                        f"{path.name}: Excel formulas are not supported (row {row_number}); "
                        "export calculated values to a plain worksheet first."
                    )
                values = [_text(cell.value) for cell in row]
                yield row_number, values[:header_length], len(values) > header_length and any(
                    values[header_length:]
                )
        finally:
            workbook.close()
    except (OSError, IndexError, BadZipFile) as exc:
        raise ValueError(f"{path.name}: cannot read XLSX rows: {exc}") from exc


def write_xlsx(path: Path, fields: list[str], rows: list[dict[str, str]],
               numeric_columns: set[str], safe_text) -> None:
    Workbook, _ = _openpyxl()
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Results")
    sheet.append(fields)
    for row in rows:
        cells = []
        for field in fields:
            value = str(row.get(field, ""))
            if field in numeric_columns and value:
                try:
                    number = float(value)
                    if not math.isfinite(number):
                        raise ValueError('non-finite number')
                    cells.append(number)
                except ValueError:
                    # Rejected rows must retain the original invalid value.
                    cells.append(safe_text(value))
            else:
                cells.append(safe_text(value))
        sheet.append(cells)
    workbook.save(path)
