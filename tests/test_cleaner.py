import csv
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from cleaner import (DEFAULT_KEY, DEFAULT_REQUIRED, DEFAULT_NUMERIC,
                     DEMO_REQUIRED, DEMO_NUMERIC, normalize_header, process)


def write_csv(path, rows):
    with path.open("w", newline="", encoding="utf-8") as fp:
        writer = csv.writer(fp)
        writer.writerows(rows)


def output_file(results, name):
    return Path(results.output_directory) / name


class CleaningTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.src = self.root / "input"
        self.src.mkdir()
        self.out = self.root / "output"

    def test_multiple_files_aliases_duplicates_invalid(self):
        write_csv(self.src / "a.csv", [
            ["ID", "Time", "Latitude", "mag"],
            ["EQ1", "2026-10-01", "35.1", "3.1"],
            ["EQ2", "2026-10-02", "12.0", "2.8"],
        ])
        write_csv(self.src / "b.csv", [
            ["event_id", "time", "latitude_deg", "Magnitude"],
            ["EQ1", "2026-10-01", "35.1", "3.1"],
            ["EQ3", "2026-10-03", "invalid", "3.8"],
            ["", "2026-10-04", "20.1", "4.2"],
            ["EQ4", "2026-10-05", "55.0", "5.0"],
        ])
        result = process(self.src, self.out, "id", "id,time,latitude", "latitude,magnitude")
        self.assertEqual((result.processed_rows, result.clean_rows,
                          result.duplicate_rows, result.rejected_rows), (6, 3, 1, 2))
        self.assertEqual(result.issues["duplicate_id"], 1)
        self.assertEqual(result.issues["invalid_number_latitude"], 1)
        self.assertEqual(result.issues["missing_id"], 1)
        self.assertTrue(output_file(result, "report.html").exists())
        with output_file(result, "cleaned.csv").open(encoding="utf-8-sig", newline="") as fp:
            rows = list(csv.DictReader(fp))
        self.assertEqual({r["id"] for r in rows}, {"EQ1", "EQ2", "EQ4"})
        self.assertEqual(rows[0]["magnitude"], "3.1")

    def test_missing_key_fails_without_outputs(self):
        write_csv(self.src / "one.csv", [["item", "name"], ["a", "Alice"]])
        with self.assertRaisesRegex(ValueError, "missing"):
            process(self.src, self.out)
        self.assertFalse(self.out.exists())

    def test_missing_configured_numeric_column_fails(self):
        write_csv(self.src / "one.csv", [["id", "amount"], ["1", "not a number"]])
        with self.assertRaisesRegex(ValueError, "price"):
            process(self.src, self.out, numeric="price")
        self.assertFalse(self.out.exists())

    def test_missing_configured_column_in_one_file_fails(self):
        write_csv(self.src / "a.csv", [["id", "price"], ["1", "10"]])
        write_csv(self.src / "b.csv", [["id", "name"], ["2", "Bob"]])
        with self.assertRaisesRegex(ValueError, "b.csv"):
            process(self.src, self.out, numeric="price")

    def test_protect_against_overlapping_paths(self):
        write_csv(self.src / "a.csv", [["id"], ["1"]])
        with self.assertRaisesRegex(ValueError, "separate"):
            process(self.src, self.src / "output")

    def test_formula_sanitization_csv(self):
        write_csv(self.src / "a.csv", [["id", "name"], ["1", '=HYPERLINK("test")']])
        result = process(self.src, self.out)
        with output_file(result, "cleaned.csv").open(encoding="utf-8-sig", newline="") as fp:
            record = next(csv.DictReader(fp))
        self.assertTrue(record["name"].startswith("'="))

    def test_formula_sanitization_xlsx_output(self):
        from openpyxl import load_workbook
        write_csv(self.src / "a.csv", [["id", "name", "price"], ["1", "=HYPERLINK(1)", "12.5"]])
        result = process(self.src, self.out, numeric="price")
        wb = load_workbook(output_file(result, "cleaned.xlsx"), read_only=True)
        try:
            row = list(wb.active.values)[1]
        finally:
            wb.close()
        self.assertEqual(row[1], "'=HYPERLINK(1)")
        self.assertEqual(row[2], 12.5)

    def test_normalize_header(self):
        self.assertEqual(normalize_header(" Latitude Deg "), "latitude")
        self.assertEqual(normalize_header("Event ID"), "id")

    def test_duplicate_headers_not_silent(self):
        write_csv(self.src / "a.csv", [["ID", "event_id"], ["one", "two"]])
        with self.assertRaisesRegex(ValueError, "duplicate"):
            process(self.src, self.out)

    def test_strict_unclosed_csv_quote(self):
        (self.src / "broken.csv").write_text('id,description\nONE,"hello\nTWO,world\n', encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "malformed CSV"):
            process(self.src, self.out)
        self.assertFalse(self.out.exists())

    def test_multiline_physical_row_tracking(self):
        write_csv(self.src / "multiline.csv", [["id", "note"], ["one", "one\ntwo"], ["two", "done"]])
        result = process(self.src, self.out)
        with output_file(result, "cleaned.csv").open(encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
        self.assertEqual([r["__source_row"] for r in rows], ["2", "4"])

    def test_mixed_csv_and_xlsx_input(self):
        from openpyxl import Workbook, load_workbook
        write_csv(self.src / "a.csv", [["ID", "Time", "Latitude"], ["A", "2026", "15.5"]])
        wb = Workbook()
        ws = wb.active
        ws.append(["event_id", "time", "latitude_deg"])
        ws.append(["B", "2026", 32.4])
        ws.append(["A", "2026", 15.5])
        wb.save(self.src / "b.xlsx")
        result = process(self.src, self.out, numeric="latitude")
        self.assertEqual((result.input_files, result.clean_rows, result.duplicate_rows), (2, 2, 1))
        workbook = load_workbook(output_file(result, "cleaned.xlsx"), read_only=True)
        try:
            values = list(workbook.active.values)
        finally:
            workbook.close()
        self.assertEqual(len(values), 3)
        self.assertEqual(values[1][2], 15.5)

    def test_excel_formulas_are_not_silently_accepted(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["id", "price"])
        ws.append(["A", "=1+2"])
        wb.save(self.src / "formula.xlsx")
        with self.assertRaisesRegex(ValueError, "formulas are not supported"):
            process(self.src, self.out, numeric="price")
        self.assertFalse(self.out.exists())

    def test_prior_output_preserved_on_failure(self):
        write_csv(self.src / "a.csv", [["id", "price"], ["A", "1"]])
        first = process(self.src, self.out, numeric="price")
        previous = output_file(first, "cleaned.csv").read_bytes()
        with patch("cleaner.write_xlsx", side_effect=OSError("simulated failure")):
            with self.assertRaisesRegex(OSError, "simulated failure"):
                process(self.src, self.out, numeric="price")
        self.assertEqual(output_file(first, "cleaned.csv").read_bytes(), previous)
        self.assertEqual(len(list(self.out.glob("run_*"))), 1)
        self.assertEqual(len(list(self.out.glob(".working_*"))), 0)

    def test_successful_runs_do_not_overwrite(self):
        write_csv(self.src / "a.csv", [["id"], ["A"]])
        first = process(self.src, self.out)
        second = process(self.src, self.out)
        self.assertNotEqual(first.output_directory, second.output_directory)
        self.assertTrue(output_file(first, "cleaned.csv").exists())
        self.assertTrue(output_file(second, "cleaned.xlsx").exists())

    def test_shared_defaults_are_generic(self):
        import gui
        self.assertEqual((DEFAULT_KEY, DEFAULT_REQUIRED, DEFAULT_NUMERIC),
                         ("id", "id", ""))
        self.assertEqual(DEMO_REQUIRED, "id,time,latitude,longitude")
        self.assertEqual(DEMO_NUMERIC, "latitude,longitude,depth,magnitude")
        self.assertIs(gui.DEFAULT_KEY, DEFAULT_KEY)

    def test_corrupt_xlsx_is_handled(self):
        (self.src / "broken.xlsx").write_bytes(b"not a real workbook")
        with self.assertRaisesRegex(ValueError, "cannot read XLSX"):
            process(self.src, self.out)
        self.assertFalse(self.out.exists())

    def test_nonfinite_numeric_values_are_rejected_in_excel_export(self):
        from openpyxl import load_workbook
        write_csv(self.src / "a.csv", [["id", "amount"], ["A", "NaN"], ["B", "Infinity"]])
        result = process(self.src, self.out, numeric="amount")
        self.assertEqual((result.clean_rows, result.rejected_rows), (0, 2))
        book = load_workbook(output_file(result, "rejected.xlsx"), read_only=True)
        try:
            rows = list(book.active.values)
        finally:
            book.close()
        self.assertEqual([rows[1][1], rows[2][1]], ["NaN", "Infinity"])

    def test_ragged_csv_rows_are_rejected(self):
        write_csv(self.src / "ragged.csv", [["id", "name"], ["A", "Alice"], ["B"], ["C", "Carol", "Extra"]])
        result = process(self.src, self.out)
        self.assertEqual((result.clean_rows, result.rejected_rows), (1, 2))
        self.assertEqual(result.issues["malformed_csv_row"], 2)


if __name__ == "__main__":
    unittest.main()
